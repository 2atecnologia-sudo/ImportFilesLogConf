import sys, json, subprocess, configparser
from pathlib import Path
from collections import defaultdict, Counter
from datetime import datetime
from PySide6.QtCore import Qt, QThread, Signal, QTimer
from PySide6.QtGui import QColor, QPainter, QIcon, QPixmap, QPen
from PySide6.QtWidgets import (QApplication,QMainWindow,QWidget,QVBoxLayout,QHBoxLayout,QLabel,QPushButton,QComboBox,QLineEdit,QTableWidget,QTableWidgetItem,QHeaderView,QFileDialog,QMessageBox,QDialog,QTabWidget,QSpinBox,QCheckBox,QGroupBox,QFormLayout,QPlainTextEdit,QScrollArea,QProgressBar,QInputDialog)
from .style import load_style
from .data import load,norm,number,document,related,format_time,order_items,item_totals
from .license import License
from .importer_config import locate
from .importer_monitor import ensure_running,open_tray
from .users import Users
from .reset_data import preview,clear
BASE=Path(sys.executable).parent if getattr(sys,'frozen',False) else Path(__file__).resolve().parents[1]

class Worker(QThread):
    done=Signal(object)
    failed=Signal(str)
    progress=Signal(str)
    def __init__(self,folder,validate=False):
        super().__init__(); self.folder=folder; self.validate=validate
    def run(self):
        try:
            self.progress.emit('Localizando a configuração do importador…')
            folder=locate(self.folder)
            self.progress.emit('Verificando a licença compartilhada…')
            license_info=License().check(folder)
            self.progress.emit("Licença encontrada. Carregando conferências…")
            self.done.emit((load(folder),license_info,str(folder)))
        except Exception as exc: self.failed.emit(str(exc))

class MonitorWorker(QThread):
    done=Signal(object)
    def __init__(self,base,folder):
        super().__init__();self.base=base;self.folder=folder
    def run(self):
        try:self.done.emit(ensure_running(self.base,self.folder))
        except Exception as exc:self.done.emit(('stopped',str(exc)))

class TrayWorker(QThread):
    done=Signal(bool,str)
    def __init__(self,base,folder):
        super().__init__();self.base=base;self.folder=folder
    def run(self):
        try:self.done.emit(True,open_tray(self.base,self.folder))
        except Exception as exc:self.done.emit(False,str(exc))

class ResetWorker(QThread):
    done=Signal(bool,object)
    def __init__(self,folder,erase=False):
        super().__init__();self.folder=folder;self.erase=erase
    def run(self):
        try:self.done.emit(True,clear(self.folder) if self.erase else preview(self.folder))
        except Exception as exc:self.done.emit(False,str(exc))

def status_icon(status):
    pixmap=QPixmap(24,24);pixmap.fill(Qt.transparent)
    painter=QPainter(pixmap);painter.setRenderHint(QPainter.Antialiasing)
    colors={'Não iniciado':'#E5B51B','Em andamento':'#EF7D22','Conferido':'#2E9D59'}
    painter.setPen(Qt.NoPen);painter.setBrush(QColor(colors.get(status,'#88939F')));painter.drawEllipse(2,2,20,20)
    pen=QPen(QColor('white'));pen.setWidth(2);pen.setCapStyle(Qt.RoundCap);pen.setJoinStyle(Qt.RoundJoin);painter.setPen(pen)
    if status=='Conferido':
        painter.drawLine(6,12,10,16);painter.drawLine(10,16,18,8)
    elif status=='Em andamento':
        painter.drawEllipse(6,6,12,12);painter.drawLine(12,8,12,12);painter.drawLine(12,12,15,14)
    else:
        painter.drawLine(9,8,9,16);painter.drawLine(15,8,15,16)
    painter.end();return QIcon(pixmap)

class MonitorIndicator(QWidget):
    def __init__(self):
        super().__init__();self.setFixedSize(18,18);self.state='stopped';self.angle=0
        self.animation=QTimer(self);self.animation.setInterval(100);self.animation.timeout.connect(self.advance)
    def set_state(self,state):
        self.state=state
        if state in ('running','starting'):self.animation.start()
        else:self.animation.stop()
        self.update()
    def advance(self):
        self.angle=(self.angle+30)%360;self.update()
    def paintEvent(self,event):
        p=QPainter(self);p.setRenderHint(QPainter.Antialiasing);p.setPen(Qt.NoPen)
        p.setBrush(QColor({'running':'#2E9D59','starting':'#EF7D22'}.get(self.state,'#D9534F')));p.drawRoundedRect(1,1,16,16,3,3)
        if self.state in ('running','starting'):
            pen=QPen(QColor('white'));pen.setWidth(2);pen.setCapStyle(Qt.RoundCap);p.setPen(pen);p.setBrush(Qt.NoBrush);p.drawArc(4,4,10,10,self.angle*16,240*16)
        p.end()

class Bars(QWidget):
    def __init__(self):
        super().__init__(); self.values=[]; self.setMinimumHeight(160)
    def paintEvent(self,event):
        p=QPainter(self); p.setRenderHint(QPainter.Antialiasing)
        if not self.values:
            p.drawText(self.rect(),Qt.AlignCenter,'Sem dados para os filtros selecionados'); return
        width=max(1,self.width()-250); maximum=max(v for _,v in self.values) or 1
        for i,(label,value) in enumerate(self.values[:8]):
            y=12+i*27
            p.setPen(QColor('#505050')); p.drawText(8,y+16,str(label)[:24])
            p.setBrush(QColor('#1976D2')); p.setPen(Qt.NoPen)
            p.drawRoundedRect(170,y,int(width*value/maximum),18,4,4)
            p.setPen(QColor('#505050')); p.drawText(180+int(width*value/maximum),y+15,f'{value:g}')
        p.end()

class Window(QMainWindow):
    def __init__(self):
        super().__init__(); self.setWindowTitle('2A • Retaguarda de Conferência'); self.resize(1280,820)
        self.heads=[];self.items=[];self.occurrences=[];self.worker=None;self.licensed=False;self.license_info=None;self.importer_folder=None;self.monitor_worker=None;self.tray_worker=None;self.tray_requested=False;self.reset_worker=None;self.resetting=False;self.last_monitor_result=None
        self.settings_path=BASE/'retaguarda.json'
        try:self.settings=json.loads(self.settings_path.read_text(encoding='utf-8'))
        except (OSError,ValueError):self.settings={'importer':r'C:\Projetos2\ImportFilesLogConf\dist','interval':30}
        central=QWidget();self.setCentralWidget(central);outer=QVBoxLayout(central);outer.setContentsMargins(12,8,12,24)
        scroll=QScrollArea();scroll.setWidgetResizable(True);scroll.setFrameShape(QScrollArea.NoFrame)
        content=QWidget();scroll.setWidget(content);outer.addWidget(scroll,1)
        layout=QVBoxLayout(content);layout.setContentsMargins(8,6,8,6);layout.setSpacing(8)
        bottom_actions=QHBoxLayout()
        top=QHBoxLayout();top.setSpacing(20);heading=QVBoxLayout();heading.setSpacing(11);title=QLabel('Retaguarda de Conferência');title.setStyleSheet('font-size:22pt;font-weight:bold;color:#1976D2');heading.addWidget(title);subtitle=QLabel('Acompanhe recebimento, expedição e separação em uma única visão.');subtitle.setWordWrap(True);subtitle.setStyleSheet('color:#1976D2;');heading.addWidget(subtitle);heading.addStretch();top.addLayout(heading,2)
        for text,fn in [('Atualizar',self.refresh),('Exportar Excel',self.export),('Configurações',self.configure),('Usuários',self.manage_users),('Zerar dados',self.reset_access),('Sobre',self.about)]:
            b=QPushButton(text);b.clicked.connect(fn);bottom_actions.addWidget(b)
        self.summary=QLabel('Aguardando a configuração e a licença do importador.');self.summary.setWordWrap(True);self.summary.setStyleSheet('background:#EEF5FC;padding:10px;border-radius:8px;font-size:13pt');top.addWidget(self.summary,3);layout.addLayout(top)
        self.loading=QWidget();loading_box=QVBoxLayout(self.loading);loading_box.setContentsMargins(10,6,10,6);loading_box.setSpacing(4)
        self.loading_text=QLabel('Iniciando a Retaguarda…');self.loading_text.setStyleSheet('color:#1976D2;font-weight:bold;');loading_box.addWidget(self.loading_text)
        self.loading_bar=QProgressBar();self.loading_bar.setRange(0,0);self.loading_bar.setTextVisible(False);self.loading_bar.setFixedHeight(16);self.loading_bar.setStyleSheet('QProgressBar {border:1px solid #D6E4F2;border-radius:4px;background:#EEF5FC;} QProgressBar::chunk {background:#1976D2;border-radius:4px;}');loading_box.addWidget(self.loading_bar);layout.addWidget(self.loading)
        filters=QHBoxLayout();self.process=QComboBox();self.process.addItems(['Todos os processos','RECEBIMENTO','EXPEDIÇÃO','SEPARAÇÃO']);self.status=QComboBox();self.status.addItems(['Todos os status','Não iniciado','Em andamento','Conferido']);self.search=QLineEdit();self.search.setPlaceholderText('Pesquisar NF, cliente, usuário ou coletor');self.auto=QCheckBox('Atualização automática');self.auto.setChecked(True)
        for widget in [self.process,self.status,self.search,self.auto]: filters.addWidget(widget)
        layout.addLayout(filters)
        self.tabs=QTabWidget();layout.addWidget(self.tabs,1)
        page=QWidget();box=QVBoxLayout(page);self.table=self.new_table();box.addWidget(self.table);actions=QHBoxLayout()
        for text,fn in [('Ver itens',self.details),('Ver ocorrências',self.show_occurrences)]:
            b=QPushButton(text);b.clicked.connect(fn);actions.addWidget(b)
        bottom_actions.insertLayout(0,actions);bottom_actions.addStretch();self.tabs.addTab(page,'Conferências')
        page=QWidget();box=QVBoxLayout(page);box.addWidget(QLabel('Quantidade lida acumulada por coletor • mesmos filtros da lista'));self.chart=Bars();box.addWidget(self.chart);self.productivity=self.new_table();box.addWidget(self.productivity);box.addWidget(QLabel('Usuário = responsável pelo início da conferência. Quantidades representam o estado atual; não são eventos de leitura.'));self.tabs.addTab(page,'Produtividade')
        monitor=QGroupBox('Monitoramento');monitor_box=QVBoxLayout(monitor);monitor_box.setContentsMargins(10,10,10,20)
        self.monitor_status=QLabel('Aguardando configuração e licença do importador.');monitor_box.addWidget(self.monitor_status)
        self.logs=QPlainTextEdit();self.logs.setReadOnly(True);self.logs.setMaximumBlockCount(200);self.logs.setFixedHeight(65);monitor_box.addWidget(self.logs)
        layout.addWidget(monitor)
        self.monitor_indicator=MonitorIndicator();self.monitor_caption=QLabel('Monitor parado');bottom_actions.addWidget(self.monitor_indicator);bottom_actions.addWidget(self.monitor_caption)
        outer.addLayout(bottom_actions)
        self.auto.toggled.connect(self.monitor_mode)
        self.table.cellClicked.connect(self.open_clicked_item)
        self.process.currentTextChanged.connect(self.render);self.status.currentTextChanged.connect(self.render);self.search.textChanged.connect(self.render)
        self.timer=QTimer(self);self.timer.timeout.connect(self.auto_refresh);self.timer.start(max(10,int(self.settings.get('interval',30)))*1000)
        self.watchdog=QTimer(self);self.watchdog.timeout.connect(self.check_monitor);self.watchdog.start(30000);QTimer.singleShot(300,self.check_monitor)
        self.statusBar().showMessage('Aguardando o importador. A configuração SQL e a licença são compartilhadas.');QTimer.singleShot(100,self.refresh)
    def call_tray(self):
        if self.resetting:return
        if self.tray_worker and self.tray_worker.isRunning():
            self.statusBar().showMessage('Aguarde: verificando o Tray…',10000);return
        self.statusBar().showMessage('Verificando e abrindo o Tray…');self.tray_worker=TrayWorker(str(BASE),self.importer_folder);self.tray_worker.done.connect(self.tray_opened);self.tray_worker.start()
        worker=self.tray_worker
        QTimer.singleShot(25000,lambda:self.tray_delayed(worker))
    def tray_delayed(self,worker):
        if self.tray_worker is worker and worker.isRunning():
            message='A verificação do Tray está demorando. Confira o processo ImportFilesLogConfTray no Gerenciador de Tarefas.';self.statusBar().showMessage(message);self.log_event(message)
    def tray_opened(self,success,message):
        self.log_event(message)
        if success:
            if not (self.worker and self.worker.isRunning()):self.statusBar().showMessage(message,15000)
        else:QMessageBox.warning(self,'Tray do importador',message)
    def check_monitor(self):
        if self.resetting:return
        if self.monitor_worker and self.monitor_worker.isRunning():return
        self.monitor_worker=MonitorWorker(str(BASE),self.importer_folder);self.monitor_worker.done.connect(self.monitor_checked);self.monitor_worker.start()
    def monitor_checked(self,result):
        state,message=result
        self.monitor_indicator.set_state(state)
        self.monitor_caption.setText({'running':'Monitor rodando','starting':'Monitor iniciando'}.get(state,'Monitor parado'))
        self.monitor_indicator.setToolTip(message);self.monitor_caption.setToolTip(message)
        if result!=self.last_monitor_result:self.log_event(message)
        self.last_monitor_result=result
    def open_clicked_item(self,row,column):
        self.table.setCurrentCell(row,column);self.details()
    def new_table(self):
        table=QTableWidget();table.setAlternatingRowColors(True);table.setSelectionBehavior(QTableWidget.SelectRows);table.setSelectionMode(QTableWidget.SingleSelection);table.setEditTriggers(QTableWidget.NoEditTriggers);table.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive);table.horizontalHeader().setStretchLastSection(True);return table
    def fill(self,table,cols,rows):
        table.setSortingEnabled(False);table.clear();table.setColumnCount(len(cols));table.setHorizontalHeaderLabels(cols);table.setRowCount(len(rows))
        for r,row in enumerate(rows):
            for c,value in enumerate(row):
                if cols[c] in ('Hora início','Hora fim','Hora erro','horainiconf','horafimconf','horaerro'):value=format_time(value)
                table.setItem(r,c,QTableWidgetItem('' if value is None else str(value)))
        table.resizeColumnsToContents()
    def auto_refresh(self):
        if self.auto.isChecked():self.refresh()
    def refresh(self,*args):
        if self.resetting:return
        if self.worker and self.worker.isRunning():return
        if not self.licensed:self.loading_text.setText('Carregando configuração, licença e conferências…');self.loading.show()
        self.log_event('Consultando configuração, licença e dados do SQL.');self.monitor_status.setText('Consultando SQL…');self.statusBar().showMessage('Consultando licença e dados…');self.worker=Worker(str(BASE),True);self.worker.done.connect(self.received);self.worker.failed.connect(self.error);self.worker.progress.connect(self.monitor_progress);self.worker.start()
    def received(self,result):
        self.loading.hide()
        (self.heads,self.items,self.occurrences),info,self.importer_folder=result;self.licensed=True
        if info:self.license_info=info
        if not self.tray_requested:
            self.tray_requested=True;QTimer.singleShot(0,self.call_tray)
        self.render();now=datetime.now();self.statusBar().showMessage('Atualizado em '+now.strftime('%d/%m/%Y %H:%M:%S'))
        self.log_event(f'Consulta concluída: {len(self.heads)} conferências, {len(self.items)} itens e {len(self.occurrences)} ocorrências.')
        if self.auto.isChecked():
            from datetime import timedelta
            next_time=(now+timedelta(milliseconds=max(0,self.timer.remainingTime()))).strftime('%H:%M:%S')
            self.monitor_status.setText('Monitorando SQL • Aguardando novas leituras • Última atualização: '+now.strftime('%H:%M:%S')+' • Próxima consulta: '+next_time)
        else:self.monitor_status.setText('Atualização automática pausada • Última consulta: '+now.strftime('%H:%M:%S'))
    def error(self,text):
        self.loading.hide()
        self.licensed=False;self.license_info=None;self.heads=[];self.items=[];self.occurrences=[];self.render();self.summary.setText('Atualização bloqueada: '+text);self.statusBar().showMessage(text);self.monitor_status.setText('Atualização bloqueada • '+text);self.log_event('ERRO: '+text)
    def log_event(self,text):
        self.logs.appendPlainText(datetime.now().strftime('%d/%m/%Y %H:%M:%S')+' | '+text)
    def monitor_progress(self,text):
        self.loading_text.setText(text)
        self.statusBar().showMessage(text);self.monitor_status.setText(text);self.log_event(text)
    def monitor_mode(self,enabled):
        text='Monitoramento automático ativado • Consulta a cada 30 segundos.' if enabled else 'Monitoramento automático pausado • Use Atualizar para consultar.'
        self.monitor_status.setText(text);self.log_event(text)
    def classify(self,h):
        status=norm(h.get('statusconf'))
        if status in ('CONFERIDO','FINALIZADO','CONCLUIDO'):return 'Conferido'
        if status in ('EM ANDAMENTO','EMANDAMENTO'):return 'Em andamento'
        return 'Não iniciado'
    def render(self,*args):
        proc=norm(self.process.currentText());query=norm(self.search.text());status=self.status.currentText()
        self.visible=[h for h in self.heads if (self.process.currentIndex()==0 or norm(h.get('processo'))==proc) and (self.status.currentIndex()==0 or self.classify(h)==status) and (not query or query in norm(' '.join(str(h.get(k) or '') for k in ['numnf','nomecli','coletorid','useriniconf','userfimconf'])))]
        counts=Counter(self.classify(h) for h in self.visible);selected_items=[i for h in self.visible for i in related(h,self.items)]
        expected=sum(number(i.get('qtdedoc')) for i in selected_items);read=sum(number(i.get('qtdelido')) for i in selected_items)
        progress=100*sum(min(max(number(i.get('qtdelido')),0),max(number(i.get('qtdedoc')),0)) for i in selected_items)/expected if expected>0 else 0
        self.summary.setText(f"{len(self.visible)} conferências   •   {counts['Não iniciado']} não iniciadas   •   {counts['Em andamento']} em andamento   •   {counts['Conferido']} conferidas\nEsperado: {expected:g}   |   Lido: {read:g}   |   Atendimento: {progress:.1f}%")
        self.fill(self.table,['Num Doc','Cliente / fornecedor','Processo','Status','Coletor','Usuário início','Usuário fim','Data conferência','Hora início','Hora fim'],[[('Não iniciado' if k=='statusconf' and norm(h.get(k))=='AGUARDANDO' else h.get(k)) for k in ['numnf','nomecli','processo','statusconf','coletorid','useriniconf','userfimconf','dataconf','horainiconf','horafimconf']] for h in self.visible])
        for row,h in enumerate(self.visible):
            item=self.table.item(row,3);state=self.classify(h);item.setIcon(status_icon(state));item.setToolTip(state)
        self.table.resizeColumnToContents(3)
        groups=defaultdict(lambda:[0,0.,0.])
        for h in self.visible:
            key=(str(h.get('coletorid') or 'Não atribuído'),str(h.get('useriniconf') or 'Não informado'));values=groups[key];values[0]+=1
            for i in related(h,self.items):values[1]+=number(i.get('qtdedoc'));values[2]+=number(i.get('qtdelido'))
        ranked=sorted(groups.items(),key=lambda kv:kv[1][2],reverse=True)
        self.fill(self.productivity,['Coletor','Usuário responsável','Conferências','Esperado','Lido','Progresso'],[[*k,*v,''] for k,v in ranked])
        for row,(_,values) in enumerate(ranked):
            expected,read=values[1],values[2];percent=read/expected*100 if expected>0 else 0
            progress=QProgressBar();progress.setRange(0,1000);progress.setValue(round(max(0,min(percent,100))*10));progress.setFormat(f'{percent:.1f}% • {read:g} / {expected:g}');progress.setAlignment(Qt.AlignCenter)
            progress.setStyleSheet('QProgressBar {border:1px solid #D6E4F2;border-radius:5px;background:#EEF5FC;text-align:center;color:#16436A;min-height:22px;} QProgressBar::chunk {background:#7AB8F2;border-radius:4px;}')
            self.productivity.setCellWidget(row,5,progress)
        self.productivity.setColumnWidth(5,280)
        collectors=defaultdict(float)
        for (collector,_),v in groups.items():collectors[collector]+=v[2]
        self.chart.values=sorted(collectors.items(),key=lambda kv:kv[1],reverse=True);self.chart.update()
    def selected(self):
        row=self.table.currentRow()
        if row<0 or row>=len(self.visible):QMessageBox.information(self,'Conferências','Selecione uma conferência.');return None
        return self.visible[row]
    def dialog_table(self,title,data,columns=None,header=None,refresh_data=None,item_summary=False):
        dlg=QDialog(self);dlg.setWindowTitle(title);dlg.resize(1100,580);box=QVBoxLayout(dlg)
        if header:
            label=QLabel(header);label.setWordWrap(True);label.setStyleSheet('font-weight:bold;font-size:12pt;color:#1976D2;padding:8px;');box.addWidget(label)
        if item_summary:
            totals_label=QLabel(item_totals(data));totals_label.setWordWrap(True);totals_label.setStyleSheet('background:#EEF5FC;padding:10px;border-radius:6px;font-weight:bold;');box.addWidget(totals_label)
            progress=QProgressBar();progress.setRange(0,1000);progress.setFixedHeight(20);progress.setAlignment(Qt.AlignCenter);progress.setStyleSheet('QProgressBar {border:1px solid #D6E4F2;border-radius:5px;background:#EEF5FC;color:#16436A;text-align:center;} QProgressBar::chunk {background:#7AB8F2;border-radius:4px;}');box.addWidget(progress)
            def update_progress(records):
                expected=sum(number(row.get('qtdedoc')) for row in records);read=sum(number(row.get('qtdelido')) for row in records);percent=read/expected*100 if expected>0 else 0
                progress.setValue(round(max(0,min(percent,100))*10));progress.setFormat(f'{percent:.1f}% conferido • {read:g} / {expected:g}')
            update_progress(data)
        table=self.new_table();box.addWidget(table)
        if columns is None:columns=[(k,k) for k in dict.fromkeys(k for r in data for k in r)]
        self.fill(table,[label for _,label in columns],[[r.get(key) for key,_ in columns] for r in data]);count_label=QLabel(f'{len(data)} registros');box.addWidget(count_label)
        for col,(_,caption) in enumerate(columns):
            if caption=='Descrição':table.horizontalHeader().setSectionResizeMode(col,QHeaderView.Stretch)
        table.horizontalHeader().setStretchLastSection(False)
        actions=QHBoxLayout()
        if refresh_data:
            update=QPushButton('Atualizar');actions.addWidget(update)
            poll=QTimer(dlg);poll.setInterval(200)
            def complete():
                if self.worker and self.worker.isRunning():return
                poll.stop();update.setEnabled(True);update.setText('Atualizar')
                current=refresh_data() if self.licensed else []
                if item_summary:
                    totals_label.setText(item_totals(current) if self.licensed else 'Atualização bloqueada.');update_progress(current)
                    if not self.licensed:progress.setFormat('Atualização bloqueada.')
                self.fill(table,[caption for _,caption in columns],[[row.get(key) for key,_ in columns] for row in current])
                count_label.setText(f'{len(current)} registros' if self.licensed else 'Atualização bloqueada. Verifique a mensagem da tela principal.')
                if header and self.licensed:
                    label.setText(header.rsplit('Itens:',1)[0]+'Itens: '+str(len(current)))
            def request_refresh():
                update.setEnabled(False);update.setText('Atualizando…');self.refresh();poll.start()
            update.clicked.connect(request_refresh);poll.timeout.connect(complete)
        b=QPushButton('Exportar Excel');b.clicked.connect(lambda:self.export_table(table));actions.addWidget(b);box.addLayout(actions);dlg.exec()
    def details(self,*args):
        h=self.selected()
        if h is not None:
            data=order_items(related(h,self.items))
            if not data:
                QMessageBox.information(self,'Itens','Nenhuma leitura para este documento.');return
            columns=[('gtin','GTIN'),('codprod','Código'),('descprod','Descrição'),('qtdedoc','Qtde esperada'),('qtdelido','Qtde lida'),('saldo','Saldo'),('localizacao','Localização'),('status','Status'),('coletorid','Coletor ID')]
            header=f"{h.get('nomecli') or ''}\nNumDoc: {h.get('numnf')}   •   Processo: {h.get('processo') or ''}   •   Itens: {len(data)}"
            self.dialog_table('Detalhes da conferência',data,columns,header,lambda:order_items(related(h,self.items)),True)
    def show_occurrences(self):
        h=self.selected()
        if h is None:return
        data=[r for r in self.occurrences if document(r.get('numnota'))==document(h.get('numnf')) and (not r.get('coletorid') or str(r.get('coletorid'))==str(h.get('coletorid')))]
        if not data:
            QMessageBox.information(self,'Ocorrências','Nenhuma ocorrência encontrada.');return
        self.dialog_table(f"Ocorrências • NF {h.get('numnf')}",data)
    def export(self):self.export_table(self.table)
    def export_table(self,table):
        name,_=QFileDialog.getSaveFileName(self,'Exportar Excel','conferencias.xlsx','Excel (*.xlsx)')
        if not name:return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font,PatternFill
            wb=Workbook();ws=wb.active;ws.title='Dados';ws.append([table.horizontalHeaderItem(c).text() for c in range(table.columnCount())])
            for row in range(table.rowCount()):ws.append([table.item(row,c).text() if table.item(row,c) else '' for c in range(table.columnCount())])
            for cell in ws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='1976D2')
            ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            for column in ws.columns:ws.column_dimensions[column[0].column_letter].width=min(50,max(14,max(len(str(c.value or '')) for c in column)+2))
            wb.save(name);QMessageBox.information(self,'Exportação','Arquivo salvo.')
        except Exception as exc:QMessageBox.warning(self,'Exportação',str(exc))
    def manage_users(self):
        login=QDialog(self);login.setWindowTitle('Acesso aos usuários');layout=QVBoxLayout(login);form=QFormLayout()
        username=QLineEdit();username.setPlaceholderText('Usuário');password=QLineEdit();password.setEchoMode(QLineEdit.Password);form.addRow('Usuário',username);form.addRow('Senha',password);layout.addLayout(form)
        enter=QPushButton('Entrar');enter.clicked.connect(login.accept);layout.addWidget(enter);password.returnPressed.connect(login.accept)
        if login.exec()!=QDialog.Accepted:return
        try:
            users=Users(BASE);actor=username.text().strip().lower();role=users.authenticate(actor,password.text())
            if not role:QMessageBox.warning(self,'Acesso','Usuário ou senha inválidos.');return
        except Exception as exc:QMessageBox.warning(self,'Usuários',str(exc));return
        dlg=QDialog(self);dlg.setWindowTitle('Usuários da Retaguarda');dlg.resize(470,300);box=QVBoxLayout(dlg)
        box.addWidget(QLabel('Conectado como '+actor+' • '+('Administrador' if role=='admin' else 'Usuário')))
        selection=QComboBox();selection.addItems(sorted(users.read()) if role=='admin' else [actor]);box.addWidget(selection)
        def ask_password():
            first,ok=QInputDialog.getText(dlg,'Senha','Nova senha:',QLineEdit.Password)
            if not ok:return None
            second,ok=QInputDialog.getText(dlg,'Senha','Confirme a nova senha:',QLineEdit.Password)
            if not ok:return None
            if not first or first!=second:QMessageBox.warning(dlg,'Senha','Informe uma senha e confirme com o mesmo valor.');return None
            return first
        def change():
            value=ask_password()
            if value is None:return
            try:users.change_password(actor,selection.currentText(),value);QMessageBox.information(dlg,'Senha','Senha alterada.')
            except Exception as exc:QMessageBox.warning(dlg,'Senha',str(exc))
        button=QPushButton('Alterar senha');button.clicked.connect(change);box.addWidget(button)
        if role=='admin':
            def add():
                name,ok=QInputDialog.getText(dlg,'Novo usuário','Nome do usuário:')
                if not ok:return
                profile,ok=QInputDialog.getItem(dlg,'Perfil','Perfil do usuário:',['Usuário','Administrador'],0,False)
                if not ok:return
                value=ask_password()
                if value is None:return
                try:
                    users.add(actor,name,value,'admin' if profile=='Administrador' else 'user');selection.clear();selection.addItems(sorted(users.read()));QMessageBox.information(dlg,'Usuários','Usuário cadastrado.')
                except Exception as exc:QMessageBox.warning(dlg,'Usuários',str(exc))
            button=QPushButton('Cadastrar usuário');button.clicked.connect(add);box.addWidget(button)
        box.addWidget(QLabel('Contas locais da Retaguarda, independentes da licença do importador.'))
        close=QPushButton('Fechar');close.clicked.connect(dlg.accept);box.addWidget(close);dlg.exec()
    def reset_access(self):
        username,ok=QInputDialog.getText(self,'Zerar dados • Administrador','Usuário administrador:')
        if not ok:return
        password,ok=QInputDialog.getText(self,'Zerar dados • Administrador','Senha:',QLineEdit.Password)
        if not ok:return
        try:
            if Users(BASE).authenticate(username,password)!='admin':
                QMessageBox.warning(self,'Zerar dados','Acesso restrito a administradores. Usuário ou senha inválidos.');return
        except Exception as exc:QMessageBox.warning(self,'Zerar dados',str(exc));return
        self.reset_app()
    def reset_app(self,dialog=None):
        if self.resetting:return
        if any(w and w.isRunning() for w in (self.worker,self.monitor_worker,self.tray_worker)):
            QMessageBox.information(dialog or self,'Zerar dados','Aguarde a consulta/verificação terminar e tente novamente.');return
        try:folder=self.importer_folder or str(locate(BASE))
        except Exception as exc:QMessageBox.warning(dialog or self,'Zerar dados',str(exc));return
        self.resetting=True;self.timer.stop();self.watchdog.stop()
        self.reset_worker=ResetWorker(folder);self.reset_worker.done.connect(lambda ok,value:self.reset_preview(ok,value,folder));self.reset_worker.start()
        if dialog:dialog.accept()
    def reset_preview(self,success,value,folder):
        if not success:self.reset_finished(False,value);return
        server,database,counts=value
        message='Servidor: '+server+'\nBanco: '+database+'\n\n'+ '\n'.join(f'{table}: {count} registros' for table,count in counts.items())+'\n\nTodos os registros dessas tabelas serão apagados.\nO importador e o Tray serão encerrados durante a limpeza.\nConfiguração, licença e usuários serão mantidos.\nArquivos e dados nos coletores permanecem; se enviados novamente, poderão voltar ao SQL.\n\nDigite ZERAR para confirmar:'
        answer,ok=QInputDialog.getText(self,'Zerar todos os dados de conferência',message)
        if not ok or answer!='ZERAR':self.resume_after_reset();return
        self.reset_worker=ResetWorker(folder,True);self.reset_worker.done.connect(self.reset_finished);self.reset_worker.start()
    def resume_after_reset(self):
        self.resetting=False;self.timer.start(max(10,int(self.settings.get('interval',30)))*1000);self.watchdog.start(30000)
    def reset_finished(self,success,value):
        self.resume_after_reset()
        if success:
            self.tray_requested=False;self.heads=[];self.items=[];self.occurrences=[];self.render();self.log_event('Dados zerados: '+str(value));QMessageBox.information(self,'Zerar dados','Dados de conferência zerados. Configuração, licença e usuários preservados.');self.check_monitor();self.refresh()
        else:self.log_event('Limpeza não concluída: '+str(value));QMessageBox.warning(self,'Zerar dados',str(value))
    def configure(self):
        dlg=QDialog(self);dlg.setWindowTitle('Conexão SQL compartilhada com o importador');dlg.resize(650,400);box=QVBoxLayout(dlg)
        box.addWidget(QLabel('A retaguarda utiliza automaticamente a conexão configurada no ImportFilesLogConf.'))
        try:
            folder=Path(self.importer_folder) if self.importer_folder else locate(BASE)
            cfg=configparser.ConfigParser(interpolation=None);cfg.read(folder/'config.ini',encoding='utf-8-sig')
            form=QFormLayout()
            for key,label in [('driver','Driver ODBC'),('server','Servidor / instância'),('port','Porta'),('database','Banco de conferências'),('user','Usuário'),('trusted_connection','Autenticação Windows')]:
                edit=QLineEdit(cfg.get('sql',key,fallback=''));edit.setReadOnly(True);form.addRow(label,edit)
            form.addRow('Senha',QLabel('Utilizada automaticamente • não exibida'))
            box.addLayout(form)
        except Exception as exc:
            message=QLabel(str(exc));message.setWordWrap(True);box.addWidget(message)
        box.addWidget(QLabel('Para alterar a conexão, use Banco de dados local no importador.'))
        close=QPushButton('Fechar');close.clicked.connect(dlg.accept);box.addWidget(close);dlg.exec()
    def about(self):
        code,period=self.license_info or ('Não validado','Atualização bloqueada')
        QMessageBox.information(self,'Sobre • 2A Tecnologia','Retaguarda de Conferência 1.0.31\n\nLicença compartilhada com o ImportFilesLogConf.\nCódigo: '+code+'\nEstado: '+period+'\n\nA ativação e a renovação são feitas no importador.\nSem licença válida, a retaguarda não atualiza os dados.')
    def closeEvent(self,event):
        if self.resetting or (self.reset_worker and self.reset_worker.isRunning()) or (self.worker and self.worker.isRunning()) or (self.monitor_worker and self.monitor_worker.isRunning()) or (self.tray_worker and self.tray_worker.isRunning()):
            self.statusBar().showMessage('Aguarde a verificação terminar para fechar.');event.ignore()
        else:event.accept()

def main():
    app=QApplication(sys.argv);app.setStyleSheet(load_style()+ '\nQTableWidget {alternate-background-color:#F4F7FA;selection-background-color:#D9EAFB;selection-color:#202020;} QHeaderView::section {background:#EEF2F7;padding:8px;border:0px;} QComboBox {padding:6px;}');w=Window();w.showMaximized();sys.exit(app.exec())
if __name__=='__main__':main()
